"""Чтение четырёх листов ЭРДБ, аудит и прямые связи МКБ → группа → УКД."""
from collections import Counter, defaultdict
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import re
import zipfile

from openpyxl import load_workbook

from .icd import canonical_code, normalize_icd

SCHEMA = {
    'sp_section': ['id', 'rus_name', 'kaz_name'],
    'sp_section_point': ['id', 'rus_name', 'kaz_name'],
    'section_point_filter': ['id', 'section_id', 'section_point_id'],
    'section_diagnoses_filter': ['id', 'section_id', 'icd10'],
}
MAX_BYTES = 50 * 1024 * 1024


def normalized_name(value: str) -> str:
    return re.sub(r'\s+', ' ', value).strip().lower().replace('ё', 'е')


def cell_text(value) -> str:
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def load_erdb(data: bytes, filename: str) -> dict:
    if not data or len(data) > MAX_BYTES:
        raise ValueError('Справочник должен быть XLSX размером до 50 МБ.')
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('Распакованный XLSX превышает 100 МБ.')
    except zipfile.BadZipFile as exc:
        raise ValueError('Файл не является корректным XLSX.') from exc
    issues = []; tables = {}

    def issue(code, message, source=None):
        issues.append({'code': code, 'message': message, 'source_ref': source})

    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise ValueError('Не удалось открыть справочник XLSX.') from exc
    try:
        missing = set(SCHEMA) - set(workbook.sheetnames)
        if missing:
            raise ValueError('В XLSX отсутствуют листы: ' + ', '.join(sorted(missing)))
        for sheet, required in SCHEMA.items():
            ws = workbook[sheet]
            if ws.max_row and ws.max_row > 100001:
                raise ValueError(f'Слишком много строк в листе {sheet}.')
            iterator = ws.iter_rows(values_only=True)
            headers = [cell_text(v).strip() for v in next(iterator, ())]
            if any(headers.count(h) != 1 for h in required):
                raise ValueError(f'Лист {sheet}: необходимы уникальные колонки {required}.')
            indexes = {name: headers.index(name) for name in required}
            records = []
            for row_no, values in enumerate(iterator, 2):
                if row_no > 100001:
                    raise ValueError(f'Слишком много строк в листе {sheet}.')
                if not any(v is not None for v in values):
                    continue
                raw = {name: cell_text(values[pos]) if pos < len(values) else '' for name, pos in indexes.items()}
                ref = {'sheet': sheet, 'row': row_no, 'columns': {name: pos+1 for name, pos in indexes.items()}}
                record = {name: value.strip() for name, value in raw.items()}
                record['raw'] = raw; record['source_ref'] = ref
                record['valid_keys'] = True
                for name, value in raw.items():
                    if value != value.strip():
                        issue('OUTER_WHITESPACE', f'{sheet}/{record["id"]}/{name}: внешние пробелы.', ref)
                    if value.startswith('=') or value in ('#REF!', '#N/A', '#VALUE!', '#DIV/0!'):
                        issue('FORMULA_OR_ERROR', f'{name}: формула или ошибка вместо исходного значения.', ref)
                        record['valid_keys'] = False
                for name in required:
                    if name in ('rus_name', 'kaz_name'):
                        record[name + '_normalized'] = normalized_name(record[name])
                        if not record[name]:
                            issue('EMPTY_NAME', f'{name}: пустое название.', ref)
                    elif not record[name]:
                        issue('EMPTY_KEY', f'{name}: пустой ключ.', ref)
                        record['valid_keys'] = False
                id_pattern = r'CH\d+' if sheet == 'sp_section' else r'CHP\d+' if sheet == 'sp_section_point' else r'\d+'
                if not re.fullmatch(id_pattern, record['id']):
                    issue('INVALID_ID', f'Некорректный ID {record["id"]!r}.', ref)
                    record['valid_keys'] = False
                if 'icd10' in record:
                    record['icd'] = normalize_icd(record['icd10'])
                    record['icd_normalized'] = canonical_code(record['icd10'])
                    if record['icd']['corrections']:
                        issue('ICD_CORRECTED', f'МКБ нормализован: {record["icd10"]} → {record["icd_normalized"]}.', ref)
                    if record['icd']['status'] != 'normalized':
                        issue('ICD_NEEDS_REVIEW', f'МКБ требует проверки по выбранной версии: {record["icd10"]}.', ref)
                    if not re.fullmatch(r'[A-Z]\d{2}(?:\.\d{1,2})?', record['icd_normalized']):
                        record['valid_keys'] = False
                records.append(record)
            tables[sheet] = records
    finally:
        workbook.close()

    for sheet, records in tables.items():
        counts = Counter(row['id'] for row in records)
        for row in records:
            if counts[row['id']] > 1:
                issue('DUPLICATE_ID', f'{sheet}: повтор ID {row["id"]}.', row['source_ref'])
                row['valid_keys'] = False
        if sheet in ('sp_section', 'sp_section_point'):
            names = defaultdict(list)
            for row in records:
                if row['rus_name_normalized']:
                    names[row['rus_name_normalized']].append(row['id'])
            for name, ids in names.items():
                if len(ids) > 1:
                    issue('DUPLICATE_NAME', f'{sheet}: {name}: {ids}. Записи не объединены.')
    groups = {r['id'] for r in tables['sp_section'] if r['valid_keys']}
    points = {r['id'] for r in tables['sp_section_point'] if r['valid_keys']}
    for sheet, target in [('section_point_filter', 'section_point_id'), ('section_diagnoses_filter', 'icd_normalized')]:
        pairs = set()
        for row in tables[sheet]:
            if row['section_id'] not in groups or (target == 'section_point_id' and row[target] not in points):
                issue('BROKEN_LINK', f'{sheet}: отсутствующий или неоднозначный родитель.', row['source_ref'])
                row['valid_keys'] = False
            pair = (row['section_id'], row.get(target))
            if pair in pairs:
                issue('DUPLICATE_LINK', f'{sheet}: повтор связи {pair}.', row['source_ref'])
            pairs.add(pair)
    point_links = [r for r in tables['section_point_filter'] if r['valid_keys']]
    diagnosis_links = [r for r in tables['section_diagnoses_filter'] if r['valid_keys']]
    for group in sorted(groups - {r['section_id'] for r in diagnosis_links}):
        issue('GROUP_WITHOUT_ICD', f'Группа {group} не привязана к МКБ.')
    for group in sorted(groups - {r['section_id'] for r in point_links}):
        issue('GROUP_WITHOUT_POINTS', f'В группе {group} нет УКД.')
    for point in sorted(points - {r['section_point_id'] for r in point_links}):
        issue('POINT_WITHOUT_GROUP', f'УКД {point} не входит в группу.')
    snapshot = {'schema_version': '1.0', 'filename': Path(filename.replace('\\', '/')).name,
                'sha256': sha256(data).hexdigest(), 'size_bytes': len(data), 'tables': tables,
                'issues': issues, 'inheritance_policy': 'exact_icd_only',
                'summary': {**{name: len(rows) for name, rows in tables.items()},
                            'unique_icd': len({r['icd_normalized'] for r in diagnosis_links}),
                            'linked_groups': len({r['section_id'] for r in diagnosis_links}),
                            'issues': len(issues), 'invalid_rows': sum(not r['valid_keys'] for rows in tables.values() for r in rows)}}
    repo = ErdbRepository(snapshot)
    snapshot['indexes'] = {'icd_to_sections': dict(repo.icd_to_sections),
                           'section_to_points': dict(repo.section_to_points),
                           'point_to_sections': dict(repo.point_to_sections),
                           'point_names': dict(repo.name_index)}
    return snapshot


class ErdbRepository:
    def __init__(self, snapshot):
        self.snapshot = snapshot
        tables = snapshot['tables']
        self.groups = {r['id']: r for r in tables['sp_section'] if r['valid_keys']}
        self.points = {r['id']: r for r in tables['sp_section_point'] if r['valid_keys']}
        self.icd_to_sections = defaultdict(list); self.section_to_points = defaultdict(list)
        self.point_to_sections = defaultdict(list); self.name_index = defaultdict(list)
        self.diagnoses = defaultdict(list); self.links = defaultdict(list)
        for row in tables['section_diagnoses_filter']:
            if row['valid_keys']:
                self.icd_to_sections[row['icd_normalized']].append(row['section_id'])
                self.diagnoses[row['icd_normalized']].append(row)
        for row in tables['section_point_filter']:
            if row['valid_keys']:
                self.section_to_points[row['section_id']].append(row['section_point_id'])
                self.point_to_sections[row['section_point_id']].append(row['section_id'])
                self.links[row['section_id']].append(row)
        for point in self.points.values():
            self.name_index[point['rus_name_normalized']].append(point['id'])

    def paths_for_icd(self, code: str) -> list[dict]:
        code = canonical_code(code)
        paths = []
        for diagnosis in self.diagnoses.get(code, []):
            for link in self.links.get(diagnosis['section_id'], []):
                group = self.groups[diagnosis['section_id']]; point = self.points[link['section_point_id']]
                paths.append({'icd': code, 'section_id': group['id'], 'section_name': group['rus_name'],
                              'point_id': point['id'], 'rus_name': point['rus_name'], 'kaz_name': point['kaz_name'],
                              'diagnosis_link_id': diagnosis['id'], 'point_link_id': link['id'],
                              'source_refs': [r['source_ref'] for r in (diagnosis, group, link, point)]})
        return paths

    def search_points(self, name: str) -> list[dict]:
        query = normalized_name(name)
        return [r for r in self.points.values() if query and query in r['rus_name_normalized']]

    def has_link(self, code: str, point_id: str) -> bool:
        return any(path['point_id'] == point_id for path in self.paths_for_icd(code))
